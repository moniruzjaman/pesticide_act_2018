#!/usr/bin/env python3
"""Fix PPTX slides — replace 68→36, ৯ অধ্যায়→৩৬টি ধারা"""
from pptx import Presentation
import os

DECKS = [
    {
        "name": "khuchra",
        "path": "/home/z/my-project/public/decks/Balainashok_Ain_2018_Khuchra_Bikreta_Guide.pptx",
    },
    {
        "name": "pictorial",
        "path": "/home/z/my-project/public/decks/Balainashok_Ain_2018_Pictorial_Guide.pptx",
    },
    {
        "name": "comprehensive",
        "path": "/home/z/my-project/public/decks/বালাইনাশক আইন, ২০১৮ — খুচরা বিক্রেতার সমন্বিত ফিল্ড গাইড ও আইনগত নির্দেশিকা.pptx",
    },
]

# Replacements: old → new
REPLACEMENTS = [
    ("৬৮টি ধারায়", "৩৬টি ধারায়"),
    ("৬৮ ধারা", "৩৬ ধারা"),
    ("৬৮টি ধারা", "৩৬টি ধারা"),
    ("৯টি অধ্যায় ও ৩৬টি ধারায়", "৩৬টি ধারায়"),  # clean up double mention
    ("৯টি অধ্যায় ও", ""),  # remove chapter count
    ("৯ অধ্যায়", "৩৬ ধারা"),  # stat card
    ("অধ্যায় · ৬৮ ধারা", "৩৬ ধারা"),  # pictorial style
    ("অধ্যায় · ৩৬ ধারা", "৩৬ ধারা"),  # cleanup
]

def fix_text(text):
    result = text
    for old, new in REPLACEMENTS:
        result = result.replace(old, new)
    # Clean up any double spaces or leading/trailing issues
    result = result.replace("  ", " ").strip()
    return result

for deck in DECKS:
    path = deck["path"]
    name = deck["name"]
    print(f"\n=== Fixing {name} ===")

    prs = Presentation(path)
    changes = 0

    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                original = para.text
                if "68" in original or "৬৮" in original or "৯ অধ্যায়" in original or "৯টি অধ্যায়" in original:
                    # Apply fix to runs (preserves formatting)
                    for run in para.runs:
                        old_text = run.text
                        new_text = fix_text(old_text)
                        if old_text != new_text:
                            run.text = new_text
                            changes += 1
                            print(f"  Slide {i}: '{old_text[:60]}' → '{new_text[:60]}'")

    # Save fixed PPTX (overwrite)
    prs.save(path)
    print(f"  Saved {path} ({changes} runs changed)")

print("\nDone. All PPTX files fixed.")
